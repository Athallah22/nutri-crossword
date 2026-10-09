import json
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_word_document():
    json_path = os.path.join(os.path.dirname(__file__), '..', 'src', 'data', 'puzzles.json')
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run('DAFTAR SOAL & KUNCI JAWABAN\nTTS GIZI: OPER PENA & MUSIK')
    run_title.bold = True
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = RGBColor(23, 23, 23)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run('Pegangan Fasilitator / Guru — Media Edukasi Gizi Seimbang SMP (5 Paket Soal Terkurasi)')
    run_sub.italic = True
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(10.5)
    run_sub.font.color.rgb = RGBColor(100, 100, 100)

    doc.add_paragraph() # Spacing

    puzzles = data.get('puzzles', [])

    for p_idx, p in enumerate(puzzles, 1):
        # Section Header
        h2 = doc.add_paragraph()
        run_h2 = h2.add_run(f'PAKET SOAL #{p_idx}: {p["title"].upper()}')
        run_h2.bold = True
        run_h2.font.name = 'Calibri'
        run_h2.font.size = Pt(12)
        run_h2.font.color.rgb = RGBColor(30, 64, 175) # Blue

        desc_p = doc.add_paragraph()
        r_desc = desc_p.add_run(f'Tema / Deskripsi: {p["description"]} ({len(p["questions"])} Soal)')
        r_desc.italic = True
        r_desc.font.name = 'Calibri'
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = RGBColor(70, 70, 70)

        # Table
        table = doc.add_table(rows=1, cols=6)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers = ['No', 'Arah', 'Kategori', 'Pertanyaan (Clue)', 'Petunjuk Ekstra (Hint)', 'Kunci Jawaban']
        col_widths = [Inches(0.4), Inches(0.8), Inches(1.1), Inches(2.3), Inches(1.7), Inches(1.1)]

        hdr_cells = table.rows[0].cells
        for idx, (cell, text) in enumerate(zip(hdr_cells, headers)):
            cell.width = col_widths[idx]
            p_cell = cell.paragraphs[0]
            p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p_cell.add_run(text)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)
            
            # Header background #1E293B
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E293B"/>')
            cell._tc.get_or_add_tcPr().append(shading)

        for row_idx, q in enumerate(p['questions']):
            row_cells = table.add_row().cells
            dir_label = 'Mendatar' if q['direction'] == 'across' else 'Menurun'
            ans_text = f"{q['answer']} ({len(q['answer'])} Huruf)"
            
            values = [
                str(q['number']),
                dir_label,
                q['category'],
                q['clue'],
                q.get('hint', '-'),
                ans_text
            ]

            is_even = (row_idx % 2 == 0)
            fill_color = "F8FAFC" if is_even else "FFFFFF"

            for idx, (cell, val) in enumerate(zip(row_cells, values)):
                cell.width = col_widths[idx]
                p_cell = cell.paragraphs[0]
                p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx in [0, 1] else WD_ALIGN_PARAGRAPH.LEFT
                run = p_cell.add_run(val)
                run.font.name = 'Calibri'
                run.font.size = Pt(9)
                
                if idx == 5:
                    run.bold = True
                    run.font.color.rgb = RGBColor(22, 101, 52) # Dark green
                
                shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
                cell._tc.get_or_add_tcPr().append(shading)

        doc.add_paragraph() # Spacing between tables

    output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'Daftar_Soal_dan_Jawaban_TTS_Gizi.docx'))
    doc.save(output_path)
    print(f'Successfully created Word document at: {output_path}')

if __name__ == '__main__':
    create_word_document()
